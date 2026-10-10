import io, hashlib
import re
import json
from pathlib import Path
from openpyxl import load_workbook
from core import FIELDS, number, norm, name_norm, unit, position_type
from storage import DATA, uid
from formula_values import FormulaValues
from docx_reader import read_tables
from materials import price_key, unit_price_column
from sections import source_heading, preserve_import_sections

IMPORT_VERSION=9

def service_reason(values,cells):
    """Exact financial labels plus row structure; never ban words in a work name."""
    identity=bool(values.get('number') or values.get('code'))
    measurement=norm(values.get('unit'))
    financial_unit=not measurement or measurement in ['%','руб','руб.','сом','сом.']
    text=norm(values.get('name')).lstrip('. -')
    code=norm(values.get('code'))
    if code in ['нр','сп','зпр','зпм','эм'] and financial_unit:
        return 'Служебное обозначение начисления / финансовой разбивки'
    if identity or not financial_unit: return ''
    # A tax label alone is not enough: numbered/code-bearing positions and
    # physical units must remain work data, including incomplete positions.
    if text in {'нсп','ндс'} and measurement in {'','%'}:
        return 'Налоговое начисление НсП / НДС; не объём работы'
    if re.match(r'^(?:нр|сп)(?:\s+от\s+зп[рм])?(?:\b|$)',text) or re.match(r'^(?:накладные расходы|сметная прибыль)(?:\s*[-:(]|$)',text):
        return 'Накладные расходы / сметная прибыль'
    if re.match(r'^(?:объем|объём|вычт\.?\s*ресурсы|добав\.?\s*ресурсы)\s*:',text):
        return 'Пояснение расчёта; количество берётся только из выбранной колонки'
    if re.match(r'^(?:к\d+\s*=|коэффициент(?:ы)?\s*[:=]|глобальные начисления\b)',text):
        return 'Коэффициенты / глобальные начисления'
    if re.match(r'^(?:итого|всего)(?:\s|,|:|$)',text): return 'Итог раздела / документа'
    labels={'зарплата рабочих','зарплата основных рабочих','зарплата механизаторов','эксплуатация машин',
            'материалы','стоимость оборудования','начисление по оборудованию','стоимость эксплуатации машин',
            'сметная заработная плата','трудозатраты осн. рабочих','трудозатраты механизаторов',
            'нормативная трудоемкость','сметная стоимость матер. с оборудованием','справочно'}
    if text in labels or re.match(r'^(?:в т\.?\s*ч\.?\s*(?:зарплата|вспомогательные|ненормирован|накладные расходы|сметная прибыль)|стоимость .+\s*-|(?:шефмонтаж(?: по оборудованию)?|в том числе|материалов|запасные части|тара и упаковка|транспортные расходы|комплектация|наценка снаба|заготовительно-складские расходы|отклонение по заработной плате|косвенные расходы|материальные ресурсы не учтенные в расценках|сдача и испытание)\s*(?:-|:|$))',text):
        return 'Расчётная финансовая строка'
    return ''

def sheet_info(doc):
    return {name:dict(rows=max((r['index'] for r in rows),default=0),
                      columns=max((len(r['cells']) for r in rows),default=0))
            for name,rows in doc.get('sheets',{}).items()}

def sheet_page(doc,sheet,start=1,limit=100):
    info=sheet_info(doc)
    if sheet not in info: raise ValueError('Лист не найден')
    start=max(1,int(start));limit=max(1,min(200,int(limit)))
    return dict(sheet=sheet,start=start,**info[sheet],items=[r for r in doc['sheets'][sheet] if start<=r['index']<start+limit])

def extraction_preview(doc,config):
    # The very same import path as Apply, on a temporary document.
    temporary=dict(doc)
    import_rows(temporary,config)
    return temporary['import_report']

def section_heading(item,mapping):
    """Recognize standalone headings inside the work table, never priced rows."""
    cells=item['cells']
    def value(field):
        col=mapping.get(field)
        return cells[col] if col is not None and col<len(cells) else ''
    if value('unit') or value('quantity'): return ''
    heading=source_heading(item,mapping)
    if heading:return heading['text']
    end=max(col for field,col in mapping.items() if col is not None and field!='unit_price')+1
    texts=[(col,text) for col,text in enumerate(cells[:end]) if str(text).strip()]
    if any(norm(text).startswith(('составил','проверил','сдал:','принял:','м.п.','(должность','должность, подпись','итого','всего')) for _,text in texts): return ''
    explicit=[(col,text) for col,text in texts if re.match(r'^(?:раздел|подраздел|глава)(?=\s|[.:№\d]|$)',norm(text))]
    if explicit:
        # Preserve supported unnumbered legacy headings, but never clear an
        # incomplete position's code/number or read a different name column.
        return next((text for col,text in explicit if col in [mapping.get('name'),mapping.get('section')]
                     and not value('number') and not value('code')), '')
    # A heading may be inside the blank half of C:D or merged over A:G.
    # Numbered/code-bearing incomplete positions must remain visible as errors.
    if len(texts)!=1: return ''
    col,text=texts[0]
    if re.fullmatch(r'[\d.,\s+-]+',text): return ''
    spans=item.get('merged_spans',[])
    merged=any(start==col and stop>=mapping['unit'] for start,stop in spans)
    if merged or not value('number') and not value('code'):
        return text
    return ''

def looks_like_work_sheet(rows):
    """Keep unfamiliar units too; a missing header must not hide work data."""
    for item in rows:
        cells=[norm(value) for value in item['cells']]
        if any(value.startswith(('составил','проверил','итого','всего')) for value in cells): continue
        names=[i for i,value in enumerate(cells) if len(value)>5 and re.search(r'[а-яa-z]',value)]
        quantities=[i for i,value in enumerate(cells) if re.fullmatch(r'-?\d+(?:[.,]\d+)?',re.sub(r'\s+','',value))]
        units=[i for i,value in enumerate(cells) if re.fullmatch(r'(?:\d+\s*)?(?:м[²³23]?|кв\.м|куб\.м|шт\.?|т\.?|кг|компл\.?|м\.п\.)',value)
               or len(value)<=30 and re.fullmatch(r'[а-яa-z][а-яa-z.²³23 -]*',value) and not re.search(r'смет|объект|адрес|заказчик|сумм|стоим|руб|сом|прайс|раздел|шифр|код|наименование',value)]
        if any(len({name,quantity,measurement})==3 for name in names for quantity in quantities for measurement in units): return True
    return False

def volume_columns(cells,role='act'):
    choices=[]
    for col,value in enumerate(cells):
        label=norm(value)
        if any(word in label for word in ['стоимост','цена','сумма','формула','руб','сом','price','cost']): continue
        if not re.search(r'^(?:кол[ -]?(?:во|ич)|количество|об[ъь]?[её]м|выполнено|фактическ|quantity|qty)',label): continue
        score=2;mode='current'
        if re.search(r'накоп|с начала',label): score+=2;mode='cumulative'
        if role=='act':
            if re.search(r'период|за месяц|за отчет|за отчёт|текущ',label): score+=8;mode='current'
            elif re.search(r'выполн|факт',label): score+=4
            if re.search(r'тендер|проект|по смете|план',label): score-=5
        elif re.search(r'тендер|проект|по смете|план',label): score+=8
        elif re.search(r'выполн|факт|накоп|с начала|период',label): score-=5
        choices.append(dict(column=col,label=str(value),score=score,mode=mode))
    return sorted(choices,key=lambda x:x['score'],reverse=True)

def layout_key(cells,role):
    return role+':'+hashlib.sha256(json.dumps([norm(v) for v in cells],ensure_ascii=False).encode()).hexdigest()

def header_mapping(cells):
    patterns={'number':r'^(?:№|номер|n[оo]?\.?|позици)', 'code':r'(?:код|шифр|обоснование|номер единичной расценки)',
              'name':r'(?:наименование|название|^работа$)', 'unit':r'^(?:ед\b|ед\.|единиц)',
              'quantity':r'^(?:кол[ -]?(?:во|ич)|количество|об[ъь]?[её]м|объем)',
              'section':r'^(?:раздел|подраздел)',
              'unit_price':r'(?:цена\s*(?:за|на)\s*(?:единиц|ед\b|ед\.)|стоимость\s*(?:за|на)\s*(?:единиц|ед\b|ед\.)|unit[ _-]*price|price\s*per\s*unit)'}
    mapping={f:None for f in FIELDS}
    for col,value in enumerate(cells):
        for field,pattern in patterns.items():
            if mapping[field] is None and re.search(pattern,norm(value)): mapping[field]=col
    return mapping

def material_project(project):
    """Read legacy prices from retained source cells, without reimport/save.

    Explicit current mappings (including an unselected price) take precedence.
    Source positions, decisions, IDs, quantities and histories are retained.
    """
    documents=[]
    for doc in project['documents']:
        if not doc.get('imported') or not doc.get('sheets'):
            documents.append(doc);continue
        config=dict(doc.get('config',{}));tables=[]
        for original in config.get('tables',[config]):
            table=dict(original);mapping=dict(table.get('mapping',{}))
            if 'unit_price' not in mapping:
                mapping['unit_price']=unit_price_column(doc['sheets'].get(table.get('sheet'),[]),table.get('first_row',int(table.get('header',0))+1))
            table['mapping']=mapping;tables.append(table)
        if config.get('tables'): config['tables']=tables
        if tables: config['mapping']=tables[0]['mapping']
        rows=[]
        for original in doc.get('rows',[]):
            row=dict(original)
            if 'unit_price' not in row:
                table=next((t for t in tables if t.get('sheet')==row.get('sheet',config.get('sheet'))),{})
                col=table.get('mapping',{}).get('unit_price');raw=row.get('raw',[])
                value=raw[col] if col is not None and col<len(raw) else ''
                row['source_values']={**row.get('source_values',{}),'unit_price':value}
                price=price_key(value);row['unit_price']=str(price) if price is not None else None
            row['position_type']=position_type(row)
            price_error='Не определена цена за единицу материала: выберите её колонку'
            if row['position_type']!='material' and price_error in row.get('errors',[]):
                row['errors']=[e for e in row['errors'] if e!=price_error]
            if row.get('kind')=='work' and row['position_type']=='material' and price_key(row.get('unit_price')) is None:
                row['errors']=list(row.get('errors',[]))
                if price_error not in row['errors']:
                    row['errors'].append(price_error)
            rows.append(row)
        documents.append(dict(doc,config=config,rows=rows))
    return dict(project,documents=documents)

def header_options(rows,role):
    """Combine printed header levels, stopping before numeric work data."""
    options=[]
    for i,item in enumerate(rows[:100]):
        cells=list(item.get('header_cells',item['cells']))
        mapping=header_mapping(cells)
        if mapping['name'] is None or mapping['unit'] is None: continue
        last=item['index'];indices=[last]
        for next_row in rows[i+1:i+3]:
            choices=volume_columns(cells,role)
            probe=dict(mapping)
            if choices: probe['quantity']=choices[0]['column']
            c=next_row['cells']
            try:
                if c[probe['name']] and c[probe['unit']] and not re.fullmatch(r'\d+',c[probe['unit']]) and header_mapping(c)['name'] is None: break
            except (IndexError,TypeError): pass
            # Only continue into another header, never into a section or footer.
            following=next_row.get('header_cells',c)
            if not any(re.search(r'^(?:по смете|план|факт|выполн|за период|за месяц|накоп|с начала|кол|об[ъь]?[её]м|\d+$)',norm(v)) for v in following): break
            if section_heading(next_row,probe): break
            for col,value in enumerate(following):
                if col>=len(cells): cells.append('')
                if value and not re.fullmatch(r'\d+',value) and norm(value) not in norm(cells[col]): cells[col]=(cells[col]+' '+value).strip()
            last=next_row['index'];indices.append(last);mapping=header_mapping(cells)
        choices=volume_columns(cells,role)
        if choices: mapping['quantity']=choices[0]['column']
        if all(mapping[f] is not None for f in ['name','unit','quantity']):
            options.append((cells,mapping,choices,last,indices))
    return options

def suggest_config(doc):
    """Read explicit column labels; never infer an act's volume semantics."""
    options=[]
    for sheet,rows in doc['sheets'].items():
        candidates=header_options(rows,doc.get('role','act'))
        for cells,mapping,choices,header,indices in candidates:
            # RIK puts standalone section headings in a blank column inside
            # the main table. Columns to the right can contain unrelated copies.
            if mapping['section'] is None:
                section_cols={}
                for entry in rows:
                    c=entry['cells']
                    if entry['index']<=header: continue
                    # Some printed name cells have an unlabelled blank half.
                    # Infer that half only between the name and unit columns;
                    # explicit manual choices never scan unrelated columns.
                    for col in range(mapping['name']+1,mapping['unit']):
                        if col not in mapping.values() and col<len(c):
                            title=section_heading(entry,dict(mapping,section=col))
                            if not title or c[col]!=title:continue
                            section_cols[col]=section_cols.get(col,0)+1
                if section_cols: mapping['section']=max(section_cols,key=section_cols.get)
            valid=0
            for item2 in rows:
                if item2['index']<=header: continue
                c=item2['cells']
                try:
                    if c[mapping['name']] and c[mapping['unit']] and not re.fullmatch(r'\d+',c[mapping['unit']]) and header_mapping(c)['name'] is None and not re_total(norm(c[mapping['name']])): valid+=1
                except (IndexError,ValueError): pass
            if valid:
                ambiguous=(len(choices)>1 and choices[0]['score']==choices[1]['score']) or bool(choices and choices[0]['mode']=='cumulative')
                header_indices=sorted({idx for option in candidates for idx in option[4]})
                options.append((valid,dict(sheet=sheet,header=header,header_indices=header_indices,mapping=mapping,section_hint=rows[0].get('section_hint',''),mode=choices[0]['mode'] if choices else 'current',method='differences',overrides={},quantity_options=choices,needs_choice=ambiguous,layout_key=layout_key(cells,doc.get('role','act')))))
        if not candidates and doc['extension']=='.docx' and options:
            previous=options[-1][1]
            width=max(len(r['cells']) for r in doc['sheets'][previous['sheet']])
            if width==max(len(r['cells']) for r in rows):
                mapping=previous['mapping'];valid=0
                for item in rows:
                    c=item['cells']
                    try:
                        if c[mapping['name']] and c[mapping['unit']]: number(c[mapping['quantity']]);valid+=1
                    except (IndexError,ValueError): pass
                if valid:
                    options.append((valid,{**previous,'sheet':sheet,'header':0,'header_indices':[], 'section_hint':rows[0].get('section_hint','') or previous.get('section_hint','')}))
    # Include every actual work sheet once; repeated printed headers are not
    # additional tables, and cover sheets without work rows are left out.
    sheets={}
    for valid,config in options:
        if config['sheet'] not in sheets or valid>sheets[config['sheet']][0]: sheets[config['sheet']]=(valid,config)
    tables=[sheets[name][1] for name in doc['sheets'] if name in sheets]
    if doc['extension'] in ['.xlsx','.xls','.docx']:
        for sheet,rows in doc['sheets'].items():
            if sheet in sheets: continue
            # A table that still looks like work data must never disappear just
            # because its labels are unfamiliar. Ask once for its columns.
            plausible=looks_like_work_sheet(rows)
            if plausible:
                tables.append(dict(sheet=sheet,header=0,mapping={f:None for f in FIELDS},mode='current',method='differences',overrides={},quantity_options=[],needs_choice=True,enabled=True,unmapped=True,layout_key=layout_key(rows[0]['cells'],doc.get('role','act'))))
    if not tables: return None
    tables.sort(key=lambda t:list(doc['sheets']).index(t['sheet']))
    first=dict(tables[0]);first['tables']=[dict(t,enabled=True) for t in tables];first['import_version']=IMPORT_VERSION
    first['needs_choice']=any(t.get('needs_choice') for t in tables)
    if len({t['mode'] for t in tables})>1: first['needs_choice']=True
    if doc['extension'] in ['.xlsx','.xls']:
        financial_reasons={'Накладные расходы / сметная прибыль','Расчётная финансовая строка','Служебное обозначение начисления / финансовой разбивки','Налоговое начисление НсП / НДС; не объём работы'}
        auxiliary=any(re.match(r'^(?:цены|текущие цены|базисные цены|начисления|определители|текущие концовки)(?:\s|$)',norm(sheet)) for sheet in doc['sheets'])
        complex_form=len(tables)>1 or auxiliary or any(service_reason({f:r['cells'][col] if col is not None and col<len(r['cells']) else '' for f,col in t['mapping'].items()},r['cells']) in financial_reasons for t in tables if t['mapping']['name'] is not None for r in doc['sheets'][t['sheet']])
        if complex_form: first['needs_choice']=True;first['manual_required']=True
    return first

def preview(content,extension):
    if extension=='.xls':
        from xls_reader import read_xls
        return read_xls(content)
    if extension=='.docx': return read_tables(content)
    if extension=='.xlsx':
        wb=load_workbook(io.BytesIO(content),data_only=True)
        formula=load_workbook(io.BytesIO(content),data_only=False)
        calculator=FormulaValues(formula,wb)
        result={}
        for sheet in wb:
            height,width=sheet.max_row,sheet.max_column
            if height>50000 or width>150: raise ValueError('Таблица превышает лимит 50 000 строк / 150 колонок')
            source_sheet=formula[sheet.title]
            rows=[]
            merged_labels={};merged_spans={}
            for merged in source_sheet.merged_cells.ranges:
                if merged.min_row==merged.max_row:
                    merged_spans.setdefault(merged.min_row,[]).append((merged.min_col-1,merged.max_col-1))
                if merged.min_row>100: continue
                label=source_sheet.cell(merged.min_row,merged.min_col).value
                for row in range(merged.min_row,min(merged.max_row,100)+1):
                    for col in range(merged.min_col,merged.max_col+1): merged_labels[(row,col)]=str(label or '')
            # Indexing a worksheet by row recalculates max_column over all cells
            # for every row. Bounded iterators keep large formatted sheets linear.
            for cells,source_cells in zip(sheet.iter_rows(max_row=height,max_col=width),source_sheet.iter_rows(max_row=height,max_col=width)):
                values=[str(c.value) if c.value is not None else '' for c in cells]
                cell_errors={str(c.column):'Ошибка ячейки Excel: '+str(c.value)
                             for c in cells if c.data_type=='e'}
                formulas=[];computed=[];formula_errors={};formula_sources={}
                for c in source_cells:
                    if c.data_type!='f': continue
                    formulas.append(c.column)
                    formula_sources[str(c.column)]=c.value
                    if values[c.column-1]: continue
                    try: values[c.column-1]=str(calculator.value(sheet.title,c.coordinate));computed.append(c.column)
                    except Exception: formula_errors[str(c.column)]='Формула требует пересчёта в Excel'
                rows.append({'index':cells[0].row,'cells':values,'header_cells':[merged_labels.get((cells[0].row,i+1),v) for i,v in enumerate(values)],'merged_spans':merged_spans.get(cells[0].row,[]),'hidden':bool(sheet.row_dimensions[cells[0].row].hidden),'formulas':formulas,'computed':computed,'formula_errors':formula_errors,'formula_sources':formula_sources,'cell_errors':cell_errors})
            result[sheet.title]=rows
        return result,[]
    if extension=='.pdf':
        import pdfplumber
        result={}; warnings=['PDF: извлечённые строки требуют проверки и обязательного подтверждения. OCR не установлен.']
        with pdfplumber.open(io.BytesIO(content)) as pdf:
            for i,page in enumerate(pdf.pages,1):
                tables=page.extract_tables()
                rows=[r for table in tables for r in table]
                if not rows: rows=[[line] for line in (page.extract_text() or '').splitlines()]
                result[f'Страница {i}']=[{'index':j,'cells':[str(v or '') for v in r],'hidden':False,'formulas':[]} for j,r in enumerate(rows,1)]
                if not rows: warnings.append(f'Страница {i}: текст не распознан. Загрузите XLSX или внесите строки вручную.')
        return result,warnings
    raise ValueError('Поддерживаются Word (.docx) и Excel (.xlsx, .xls)')

def upload(p,name,content,role):
    if role not in ['tender','act']: raise ValueError('Неверный тип документа')
    ext=Path(name).suffix.lower(); sheets,warnings=preview(content,ext)
    digest=hashlib.sha256(content).hexdigest()
    duplicates=[d for d in p['documents'] if d['hash']==digest and d['role']==role]
    if duplicates: warnings.append('Повторная загрузка одинакового файла: подтвердите включение документа')
    doc={'id':uid(),'name':Path(name).name,'role':role,'hash':digest,'extension':ext,'sheets':sheets,'warnings':warnings,'active':not bool(duplicates),'duplicate':bool(duplicates),'imported':False,'mode':'current','method':'differences','rows':[]}
    folder=DATA/p['id']; folder.mkdir(parents=True,exist_ok=True)
    (folder/(doc['id']+ext)).write_bytes(content)
    doc['suggestion']=suggest_config(doc)
    p['documents'].append(doc); return doc

def import_rows(doc,config):
    tables=config.get('tables')
    if not tables:
        _import_table(doc,config);doc['import_version']=IMPORT_VERSION;return
    active=[t for t in tables if t.get('enabled',True)]
    if not active: raise ValueError('Выберите хотя бы один лист с работами')
    selected_sheets=[t['sheet'] for t in active]
    if len(selected_sheets)!=len(set(selected_sheets)):
        raise ValueError('Один лист указан для импорта несколько раз. Выберите каждый лист только один раз.')
    rows=[];skipped=[];carry_section='';carry_context='';carry_context_name='';previous_hint=None
    for index,table in enumerate(active):
        local={**config,**table};local.pop('tables',None)
        hint=table.get('section_hint','')
        if doc['extension']=='.docx' and (not hint or hint==previous_hint):
            local['section_hint']=carry_section or hint;local['context_hint']=carry_context;local['context_name_hint']=carry_context_name
        # Single-sheet saved corrections retain their old keys.
        local['overrides']={**table.get('overrides',{}),**config.get('overrides',{})}
        temporary=dict(doc);_import_table(temporary,local,multiple=len(tables)>1,primary=table['sheet']==config['sheet'])
        rows.extend(temporary['rows'])
        skipped.extend(temporary['import_report']['skipped'])
        if temporary['rows']:
            last=temporary['rows'][-1];carry_section=last['section'];carry_context=last['context'];carry_context_name=last.get('context_name','')
            if last['kind']=='section': carry_context=carry_context_name=''
            elif last['kind']=='work' and position_type(last)=='work':
                carry_context=norm(last['code']);carry_context_name=name_norm(last['name'])
        previous_hint=hint
    # Repeated source positions are additive quantities, not duplicate rows.
    doc.update(config=config,rows=rows,imported=True,import_version=IMPORT_VERSION,mode=config.get('mode','current'),method=config.get('method','differences'),act_number=config.get('act_number',''),date=config.get('date',''))
    preserve_import_sections(doc)
    doc['import_report']=make_report(rows,skipped)

def make_report(rows,skipped):
    positions=[r for r in rows if r['kind']=='work']
    review=[r for r in rows if r['errors']]
    return dict(positions=positions,skipped=skipped,review=review,
                counts=dict(positions=len(positions),skipped=len(skipped),review=len(review)))

def _import_table(doc,config,multiple=False,primary=True):
    if doc['extension']=='.pdf' and not config.get('pdf_confirmed'): raise ValueError('Подтвердите проверку PDF')
    selected=config['sheet']; raw=doc['sheets'][selected]; mapping=config['mapping']
    if 'unit_price' not in mapping:
        mapping=dict(mapping,unit_price=unit_price_column(raw,config.get('first_row',int(config.get('header',1))+1)))
        config['mapping']=mapping
    if any(mapping.get(f) is None for f in ['name','unit','quantity']): raise ValueError('Укажите колонки наименования, единицы и объёма')
    required=[mapping[f] for f in ['name','unit','quantity']]
    if len(set(required))!=3 or any(not isinstance(col,int) or col<0 for col in required): raise ValueError('Наименование, единица и объём должны находиться в разных столбцах. Для объединённых полей нужен отдельный разбор формы.')
    assigned=[v for v in mapping.values() if v is not None]
    width=max((len(r['cells']) for r in raw),default=0)
    if any(not isinstance(v,int) or v<0 or v>=width for v in assigned): raise ValueError('Колонка находится за пределами листа')
    if len(assigned)!=len(set(assigned)): raise ValueError('Назначьте каждому полю отдельную колонку. Если шифр, название и единица объединены в ячейке, нужен отдельный разбор этой формы; поля догадками не разделяются.')
    try:
        def row_number(value):
            if not re.fullmatch(r'\d+',str(value)): raise ValueError()
            return int(value)
        first=row_number(config['first_row']) if 'first_row' in config else row_number(config.get('header',1))+1
        last=row_number(config['last_row']) if config.get('last_row') not in [None,''] else max((r['index'] for r in raw),default=0)
    except (ValueError,TypeError): raise ValueError('Укажите целые номера строк диапазона')
    if first<1 or last<first or last>max((r['index'] for r in raw),default=0): raise ValueError('Неверный диапазон строк листа')
    if config.get('mode','current') not in ['current','cumulative']: raise ValueError('Неверный режим объёмов')
    if config.get('method','differences') not in ['differences','last']: raise ValueError('Неверный способ накопления')
    rows=[];skipped=[]; section=config.get('section_hint',''); context=config.get('context_hint','');context_name=config.get('context_name_hint','')
    overrides=config.get('overrides',{})
    def skip(item,reason):
        skipped.append(dict(sheet=selected,source_row=item['index'],source=selected+':'+str(item['index']),raw=item['cells'],reason=reason))
    for item in raw:
        if item['index']<first or item['index']>last:
            if item['index']<first:
                title=section_heading(item,mapping)
                if title: section=title
            skip(item,'Вне выбранного диапазона');continue
        if item['index'] in config.get('header_indices',[]): skip(item,'Строка шапки');continue
        cells=item['cells']
        if not any(cells): skip(item,'Пустая строка');continue
        printed=header_mapping(cells)
        if all(printed[f] is not None for f in ['name','unit','quantity']): skip(item,'Повторная шапка');continue
        main=[cells[mapping[f]] if mapping[f]<len(cells) else '' for f in ['name','unit','quantity']]
        if all(re.fullmatch(r'\d+',str(v)) for v in main) and any(header_mapping(r['cells'])['name'] is not None for r in raw[max(0,item['index']-3):item['index']-1]): skip(item,'Нумерация колонок');continue
        r={f:(cells[int(mapping[f])] if mapping.get(f) is not None and int(mapping[f])<len(cells) else '') for f in FIELDS}
        source_values=dict(r)
        key=selected+':'+str(item['index'])
        override=overrides.get(key,overrides.get(str(item['index']),{}) if primary else {})
        r.update({k:v for k,v in override.items() if k in FIELDS})
        # A vertically merged quantity belongs to its first row only.
        if not override and mapping['quantity'] in item.get('continued_columns',[]): skip(item,'Продолжение объединённой ячейки; объём учитывается один раз');continue
        # Empty primary columns with auxiliary text elsewhere are not work rows.
        footer=any(re.match(r'^(?:составил|проверил|сдал:|принял:|м\.п\.|\(должность, подпись)',norm(c)) for c in cells)
        reason=service_reason(r,cells) if not override else ''
        section_title=section_heading(item,mapping) if not override else ''
        if not override and not footer and not any(v for f,v in r.items() if f!='unit_price') and not section_title: skip(item,'Нет данных в назначенных колонках');continue
        row_id=doc['id']+(':'+selected if multiple and not primary else '')+':'+str(item['index'])
        r.update(id=row_id,source=key,sheet=selected,source_row=item['index'],import_key=key,raw=cells,source_values=source_values,hidden=item['hidden'],errors=[])
        for field,col in mapping.items():
            if field=='unit_price' and position_type(r)!='material': continue
            error=item.get('cell_errors',{}).get(str(col+1)) if col is not None else None
            if error and field not in override: r['errors'].append(error)
        text=norm(r['name'])
        total_label=any(re_total(norm(cells[col])) for col in {mapping.get('name'),mapping.get('section'),mapping.get('number')} if col is not None and col<len(cells))
        # Totals have financial structure. A work named "Итого..." or
        # "Итоговая проверка..." must not disappear because of its title.
        total=bool(total_label and not r['code'] and
                   norm(r['unit']) in {'','%','руб','руб.','сом','сом.'} and
                   (not r['number'] or re_total(norm(r['number']))))
        if section_title and not r['quantity'] and not r['unit']:
            r.update(name=section_title,number='',code='')
        inferred='ignored' if footer else 'total' if total else 'ignored' if reason else 'section' if (r['name'] or r['section']) and not r['unit'] and not r['quantity'] and not r['code'] and not r['number'] else 'work'
        if inferred=='section' and config.get('manual') and not section_title: inferred='work'
        r['kind']=override.get('kind',inferred)
        if r['kind'] not in ['work','section','total','ignored']: raise ValueError('Неверный тип строки')
        if r['kind']=='section':
            r['name']=r['name'] or r['section']; context=context_name=''
            explicit=re.match(r'^(?:раздел|подраздел|глава)(?=\s|[.:№\d]|$)',norm(r['name']))
            merged=any(a<=mapping['name']<=b and b>=mapping['unit'] for a,b in item.get('merged_spans',[]))
            if not config.get('manual') or explicit or merged: section=r['name']
            if config.get('manual') and not explicit and not merged and override.get('kind')!='section':
                r['errors'].append('Одиночная текстовая строка: проверьте предполагаемый заголовок раздела; возможна неполная позиция')
        r['section']=r['section'] or section
        r['context']=context
        r['context_name']=context_name
        if r['kind']!='work':
            r['skip_reason']=reason or ('Подпись' if footer else 'Заголовок раздела' if r['kind']=='section' else 'Ручное исключение строки' if r['kind']=='ignored' and override else 'Итог раздела / документа')
            skip(item,r['skip_reason'])
        if r['kind']=='work':
            if mapping['name'] in item.get('continued_columns',[]): r['merged_name_continuation']=True
            if not r['name']: r['errors'].append('Пустое наименование')
            if not r['unit']: r['errors'].append('Пустая единица измерения')
            try:
                q=number(r['quantity']); r['quantity']=str(q)
                confirmed=config.get('negative_confirmed_rows',[])
                if doc.get('role')!='tender' and q<0 and key not in confirmed and not (primary and str(item['index']) in confirmed): r['errors'].append('Отрицательный объём требует подтверждения')
                if doc.get('role')!='tender' and abs(q)>100000000: r['errors'].append('Подозрительно большой объём: проверьте исходные данные')
            except ValueError as e: r['errors'].append(str(e)); r['quantity']=None
            r['position_type']=position_type(r)
            price=price_key(r.get('unit_price'))
            r['unit_price']=str(price) if price is not None else None
            if r['position_type']=='material' and price is None:
                r['errors'].append('Не определена цена за единицу материала: выберите её колонку')
            if r['position_type']=='work': r['context']=r['context_name']=''
            # Every repeated row contributes; only repeated FILES are blocked.
            if mapping['quantity']+1 in item.get('formulas',[]) and r['quantity'] is None: r['errors'].append('Формула без сохранённого значения: пересчитайте XLSX в Excel')
            if mapping['quantity']+1 in item.get('computed',[]): r['quantity_computed']=True
            if str(mapping['quantity']+1) in item.get('formula_sources',{}): r['quantity_formula']=item['formula_sources'][str(mapping['quantity']+1)]
            if r['position_type']=='work': context=norm(r['code']);context_name=name_norm(r['name'])
        rows.append(r)
    # A negative resource followed by its positive replacement is a normal
    # estimate correction, not a failed numeric import. Preserve both signs.
    works=[r for r in rows if r['kind']=='work']
    for old,new in zip(works,works[1:]):
        if old['quantity'] is None or new['quantity'] is None: continue
        a,b=number(old['quantity']),number(new['quantity'])
        same_unit=re.sub(r'\s+','',unit(old['unit'],True))==re.sub(r'\s+','',unit(new['unit'],True))
        old_name,new_name=name_norm(old['name']),name_norm(new['name'])
        if a<0 and b==-a and same_unit and 'прайс' in norm(new['code']) and old_name and old_name in new_name and old['section']==new['section']:
            old['errors']=[e for e in old['errors'] if 'Отрицательный' not in e]
            old['signed_adjustment']=True
    doc.update(config=config,rows=rows,imported=True,mode=config.get('mode','current'),method=config.get('method','differences'),act_number=config.get('act_number',''),date=config.get('date',''))
    preserve_import_sections(doc)
    doc['import_report']=make_report(rows,skipped)
    if not rows: raise ValueError('Нет строк данных после указанной шапки')

def re_total(text):
    return bool(re.match(r'^(?:итого|всего|накладные расходы|сметная прибыль)(?=\s|[,:;.(-]|$)',text.lstrip('. -')))

